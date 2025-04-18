from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.conf import settings
from .forms import UserLoginForm, UserRegistrationForm, SheetCreationForm
from .models import Sheet, Topic, SubTopic, Prompt, SavedSheet, LikedSheet
from django.contrib import messages
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from docx import Document
from reportlab.platypus import SimpleDocTemplate
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, Spacer
from .worksheet_utils.generation import generate_worksheet_content
import os
from .worksheet_utils.file_utils import create_sheet
# Create your views here.

def home(request):
    return render(request, "home.html")

def login(request):
    # If POST request, process form data
    if request.method == 'POST':
        form = UserLoginForm(request, data=request.POST)
        if form.is_valid():
            username = form.cleaned_data.get('username')
            password = form.cleaned_data.get('password')
            user = authenticate(username=username, password=password)
            if user is not None:
                auth_login(request, user)
                return redirect('home')
    else:
        # If GET request, create empty form
        form = UserLoginForm()
    
    return render(request, "login.html", {'form': form})

def signup(request):
    # If POST request, process form data
    if request.method == 'POST':
        form = UserRegistrationForm(request.POST)
        if form.is_valid():
            user = form.save()
            # Log the user in after registration
            auth_login(request, user)
            messages.success(request, 'Registration successful!')
            return redirect('home')
    else:
        # If GET request, create empty form
        form = UserRegistrationForm()
    
    return render(request, "signup.html", {'form': form})

def logout(request):

    if request.method == 'POST':
        auth_logout(request)
        messages.success(request, 'You have been logged out successfully.')
    return redirect('home')

@login_required
def createsheet(request):
    if request.method == 'POST':
        print("Form submitted")  # Debug print
        form = SheetCreationForm(request.POST)
        if form.is_valid():
            print("Form is valid")  # Debug print
            is_published = form.cleaned_data['published']
            # Create the sheet
            sheet = Sheet.objects.create(
                user=request.user,
                title=form.cleaned_data['title'],
                published=is_published,
                grade_level=form.cleaned_data['grade_level'],
                subject=form.cleaned_data['subject'],
                topic=form.cleaned_data['topic'],
                sub_topic=form.cleaned_data['subtopic'],
                true_false_count=int(form.cleaned_data['true_false_questions']),
                fill_in_the_blank_count=int(form.cleaned_data['fill_blank_questions']),
                multiple_choice_count=int(form.cleaned_data['multiple_choice_questions']),
                short_answer_count=int(form.cleaned_data['short_answer_questions']),
                include_answer_key=form.cleaned_data['include_answer_sheet'],
                prompt=Prompt.objects.get(type="GENERATE")
            )
            print(f"Sheet created with ID: {sheet.id}")  # Debug print
            return redirect('viewsheet', sheet_id=sheet.id)
        else:
            print("Form errors:", form.errors)  # Debug print
    else:
        form = SheetCreationForm()
    
    return render(request, 'createsheet.html', {'form': form})

def get_topics(request):
    grade_level_id = request.GET.get('grade_level')
    subject_id = request.GET.get('subject')
    
    topics = Topic.objects.filter(grade_level_id=grade_level_id, subject_id=subject_id)
    return JsonResponse(list(topics.values('id', 'name')), safe=False)

def get_subtopics(request):
    topic_id = request.GET.get('topic')
    subtopics = SubTopic.objects.filter(topic_id=topic_id)
    return JsonResponse(list(subtopics.values('id', 'name')), safe=False)

# Create sheet view and function
@login_required
def create_sheet(request):
    if request.method == 'POST':
        print("Form submitted")  # Debug print
        form = SheetCreationForm(request.POST)
        if form.is_valid():
            print("Form is valid")  # Debug print
            is_published = form.cleaned_data['published']
            # Create the sheet
            sheet = Sheet.objects.create(
                creator=request.user,
                title=form.cleaned_data['title'],
                published=is_published,
                grade_level=form.cleaned_data['grade_level'],
                subject=form.cleaned_data['subject'],
                topic=form.cleaned_data['topic'],
                subtopic=form.cleaned_data['subtopic'],
                true_false_count=int(form.cleaned_data['true_false_questions']),
                fill_blank_count=int(form.cleaned_data['fill_blank_questions']),
                multiple_choice_count=int(form.cleaned_data['multiple_choice_questions']),
                short_answer_count=int(form.cleaned_data['short_answer_questions']),
                include_answer_sheet=form.cleaned_data['include_answer_sheet']
            )
            print(f"Sheet created with ID: {sheet.id}")  # Debug print
            return redirect('view_worksheet', sheet_id=sheet.id)
        else:
            print("Form errors:", form.errors)  # Debug print
    else:
        form = SheetCreationForm()
    
    return render(request, 'create_sheet.html', {'form': form})

#View worksheet view 
@login_required
def viewsheet(request, sheet_id):
    sheet = get_object_or_404(Sheet, id=sheet_id)
    
    # Only show permission error if trying to access private sheet of another user thats not published
    if not sheet.published and sheet.user != request.user:
        messages.error(request, 'You do not have permission to view this private worksheet.')
        return redirect('home')

    # Initialize worksheet_content
    worksheet_content = ""
    
    if not sheet.docx_file or not sheet.pdf_file:
        try:
            # Only allow generation if user is the creator
            if sheet.user != request.user:
                messages.error(request, 'Only the creator can generate this worksheet.')
                return redirect('home')

            # Create the worksheet and files using the utility function
            from .worksheet_utils.file_utils import create_sheet as create_sheet_files
            create_sheet_files(sheet)
            
            # Refresh the sheet object to get updated file fields
            sheet.refresh_from_db()
            
            if not sheet.docx_file or not sheet.pdf_file:
                messages.error(request, 'Failed to generate worksheet files. Please try again.')
                return redirect('createsheet')

        except Exception as e:
            import traceback
            traceback.print_exc()
            messages.error(request, f'Error generating worksheet: {str(e)}')
            return redirect('home')

    # Check if the worksheet is saved by the current user
    is_saved = SavedSheet.objects.filter(user=request.user, sheet=sheet).exists()
    is_liked = LikedSheet.objects.filter(user=request.user, sheet=sheet).exists()

    context = {
        'sheet': sheet,
        'is_creator': sheet.user == request.user,
        'is_saved': is_saved,
        'is_liked': is_liked
    }
    
    return render(request, 'viewsheet.html', context)

@require_POST
@login_required
def toggle_like(request, sheet_id):
    sheet = get_object_or_404(Sheet, id=sheet_id)
    like, created = LikedSheet.objects.get_or_create(user=request.user, sheet=sheet)
    
    if not created:
        # User already liked the sheet, so unlike it
        like.delete()
        liked = False
    else:
        liked = True

    return JsonResponse({
        'liked': liked,
        'like_count': sheet.like_count
    })